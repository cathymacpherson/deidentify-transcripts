from __future__ import annotations

import json
import re
from datetime import date, datetime, time
from pathlib import Path

from .schemas import DeidReport, Transcript, Turn

_SPREADSHEET_REQUIRED_COLUMNS = {"speaker", "transcript"}


def load_transcript(path: Path, transcript_id: str | None = None) -> Transcript:
    suffix = path.suffix.lower()

    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if transcript_id:
            data["transcript_id"] = transcript_id
        turns = [
            Turn(
                turn_id=index,
                speaker=turn.get("speaker", "unknown"),
                text=turn.get("text", turn.get("raw_text", "")),
            )
            for index, turn in enumerate(data["turns"])
        ]
        return Transcript(
            transcript_id=data.get("transcript_id", transcript_id or path.stem),
            turns=turns,
        )

    if suffix in (".xlsx", ".xls"):
        return _load_spreadsheet_transcript(path, transcript_id)

    turns: list[Turn] = []
    for raw_line in path.read_text(encoding="utf-8-sig").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        match = re.match(r"^([^:]{1,80}):\s*(.*)$", line)
        speaker, text = (match.group(1).strip(), match.group(2)) if match else ("unknown", line)
        turns.append(Turn(turn_id=len(turns), speaker=speaker, text=text))
    return Transcript(transcript_id=transcript_id or path.stem, turns=turns)


def _read_spreadsheet_rows(path: Path) -> list[list[object]]:
    """Return every row (including the header) as a list of cell values. Date/time-formatted
    cells come back as datetime.date/time/datetime objects; everything else as-is."""
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=True, read_only=True)
        sheet = workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]

    import xlrd

    workbook = xlrd.open_workbook(str(path))
    sheet = workbook.sheet_by_index(0)
    rows: list[list[object]] = []
    for row_index in range(sheet.nrows):
        row: list[object] = []
        for cell in sheet.row(row_index):
            if cell.ctype == xlrd.XL_CELL_DATE:
                converted = xlrd.xldate.xldate_as_datetime(cell.value, workbook.datemode)
                # xlrd always returns a full datetime, anchored to its epoch date, even for a
                # cell formatted as time-only. A value under 1 day means no real date component
                # was encoded, so drop the synthetic epoch date to match openpyxl's behavior.
                if 0 <= cell.value < 1:
                    converted = converted.time()
                row.append(converted)
            elif cell.ctype == xlrd.XL_CELL_EMPTY:
                row.append(None)
            else:
                row.append(cell.value)
        rows.append(row)
    return rows


def _cell_to_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def _cell_to_time_str(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, (datetime, time, date)):
        return value.isoformat()
    text = _cell_to_text(value)
    return text or None


def _load_spreadsheet_transcript(path: Path, transcript_id: str | None) -> Transcript:
    rows = _read_spreadsheet_rows(path)
    if not rows:
        return Transcript(transcript_id=transcript_id or path.stem, turns=[])

    header = [_cell_to_text(cell).lower() for cell in rows[0]]
    missing = _SPREADSHEET_REQUIRED_COLUMNS - set(header)
    if missing:
        raise ValueError(f"spreadsheet is missing required column(s): {', '.join(sorted(missing))}")
    columns = {name: index for index, name in enumerate(header) if name}

    def cell(row: list[object], column_name: str) -> object:
        index = columns.get(column_name)
        return row[index] if index is not None and index < len(row) else None

    turns: list[Turn] = []
    for row in rows[1:]:
        if not row:
            continue
        text = _cell_to_text(cell(row, "transcript"))
        if not text:
            continue
        speaker = _cell_to_text(cell(row, "speaker")) or "unknown"
        turns.append(
            Turn(
                turn_id=len(turns),
                speaker=speaker,
                text=text,
                start_time=_cell_to_time_str(cell(row, "start time")),
                stop_time=_cell_to_time_str(cell(row, "stop time")),
            )
        )
    return Transcript(transcript_id=transcript_id or path.stem, turns=turns)


def _anonymised_turn(turn: Turn) -> dict[str, object]:
    data: dict[str, object] = {
        "turn_id": turn.turn_id,
        "speaker": turn.speaker,
        "text": turn.anonymised_text,
    }
    if turn.start_time is not None:
        data["start_time"] = turn.start_time
    if turn.stop_time is not None:
        data["stop_time"] = turn.stop_time
    return data


def save_outputs(
    transcript: Transcript,
    report: DeidReport,
    output_dir: Path,
) -> tuple[Path, Path, Path]:
    anonymised_dir = output_dir / "anonymised"
    sensitive_dir = output_dir / "sensitive"
    anonymised_dir.mkdir(parents=True, exist_ok=True)
    sensitive_dir.mkdir(parents=True, exist_ok=True)

    anonymised_path = anonymised_dir / f"{transcript.transcript_id}.json"
    anonymised_path.write_text(
        json.dumps(
            {
                "transcript_id": transcript.transcript_id,
                "turns": [_anonymised_turn(turn) for turn in transcript.turns],
            },
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    report_path = sensitive_dir / f"{transcript.transcript_id}.deid-report.json"
    report_path.write_text(report.model_dump_json(indent=2), encoding="utf-8")

    queue_path = sensitive_dir / f"{transcript.transcript_id}.review-queue.jsonl"
    with queue_path.open("w", encoding="utf-8") as handle:
        for item in report.review_items:
            handle.write(item.model_dump_json() + "\n")

    return anonymised_path, report_path, queue_path

